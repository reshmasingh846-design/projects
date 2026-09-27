#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
=====================================================================================
 60-MINUTE-AHEAD CASH-FLOW FORECASTER  -  HYBRID  CHRONOS-2 (covariates)  +  MONTE CARLO
 with calendar features, lumpy-payment split and validation-tuned error reduction
=====================================================================================
Input : minute-bucket file (CSV/XLSX):  Value_date | bucket | credit_amount | debit_amount
        optional scheduled-payments file (CSV): datetime | credit_amount | debit_amount
Target: CREDIT, DEBIT and NET total in the next 60 minutes (clock hour HH:00-HH:59)

LUMPY-PAYMENT HANDLING
  Every minute amount is classified (thresholds learned on TRAIN only) as
    * scheduled  - matches the optional scheduled-payments file   -> added as known amount
    * large      - >= 99th pct of train amounts, or a round multiple of 1,000,000
    * routine    - everything else
  TOTAL forecast = routine forecast + large-payment forecast + scheduled amount

ROUTINE FLOW  (smooth part)  - hybrid blend, weights learned on VALIDATION
  1. Chronos-2 (Amazon, zero-shot) on log1p(routine flow), credit & debit jointly
       - future-known covariates: calendar features
       - past-only covariates: large-payment history
       - context-length ensemble (default 512 / 2048 hours)
       - output statistic (mean or a quantile) chosen on validation to minimise WAPE
  2. Monte Carlo structural: compound-Poisson (count ~ Poisson, amounts bootstrapped
     from the same hour-of-week) with calendar multipliers.
  3. (optional, --use-gbm) gradient boosting as a third component / benchmark.

ERROR REDUCTION (all fitted on VALIDATION, confirmed on TEST)
  lumpy split | log scale | context ensemble | quantile selection | covariates |
  blend weights | hour-of-day calibration (shrunk) | large-payment scale

LARGE PAYMENTS  - Monte Carlo "probability x size":
     N_large ~ Poisson(rate for this hour-of-day / day-type x calendar factor),
     sizes bootstrapped from recent large payments. Gives expected value and P90.

UNCERTAINTY  - Monte Carlo joint residual bootstrap -> P10/P50/P90 (credit, debit, net).

LEAKAGE CONTROL
  * features / contexts use data <= h-1-latency only; calendar + schedule are known ahead
  * thresholds, operating hours, calendar factors: TRAIN only
  * hyper-parameters, blend weights, large-payment scale, residuals: VALIDATION only
  * TEST scored once; walk-forward refits only on already-known hours
  * automated tamper test (future values randomised -> past inputs must not change)

RUN
  pip install pandas numpy scikit-learn matplotlib openpyxl lightgbm
  pip install "chronos-forecasting>=2.0" "pandas[pyarrow]" torch   # Chronos-2 (required)
  python cashflow_forecast.py                      # uses the default paths below
  python cashflow_forecast.py --holiday-country IN --fm-device cuda
  default input : /home/singres/src/fs_gtsy_mde/src/tsfm/Trial/chaps_intraday.xlsx
  default output: /home/singres/src/fs_gtsy_mde/src/tsfm/Trial/output/
         [--schedule-file scheduled.csv] [--use-gbm] [--fm-contexts 512,2048] [--fm-device cuda]
=====================================================================================
"""
from __future__ import annotations

import argparse
import base64
import io
import itertools
import json
import time
import warnings
from dataclasses import dataclass, asdict
from pathlib import Path

import numpy as np
import pandas as pd
import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt  # noqa: E402
from sklearn.ensemble import HistGradientBoostingRegressor  # noqa: E402
from sklearn.inspection import permutation_importance  # noqa: E402

warnings.filterwarnings("ignore")
try:
    import lightgbm as lgb
except ImportError:
    lgb = None
try:
    import xgboost as xgb
except ImportError:
    xgb = None
try:
    import holidays as pyholidays
except ImportError:
    pyholidays = None

SERIES = ("credit", "debit")
HOUR = pd.Timedelta(hours=1)
Q_LEVELS = (10, 50, 90)


# ============================================================================ config
@dataclass
class Config:
    input: str = "/home/singres/src/fs_gtsy_mde/src/tsfm/Trial/chaps_intraday.xlsx"
    outdir: str = "/home/singres/src/fs_gtsy_mde/src/tsfm/Trial/output"
    sheet: str | None = None            # default: "Sheet2" if present, else first sheet
    schedule_file: str | None = None
    holiday_country: str | None = None
    holiday_file: str | None = None
    val_days: int = 90
    test_days: int = 90
    warmup_hours: int = 24 * 29
    latency_hours: int = 0
    refit_every_days: int = 30
    fm_backend: str = "chronos"         # chronos | auto (continue without) | none
    fm_model: str = "amazon/chronos-2"
    fm_device: str = "auto"             # auto | cuda | cpu
    fm_contexts: str = "512,2048"       # context-length ensemble (hours)
    fm_context_hours: int = 2048        # longest context (set automatically)
    fm_batch: int = 64
    use_gbm: bool = False               # add gradient boosting as 3rd component
    gbm_backend: str = "auto"           # auto | lightgbm | xgboost | sklearn
    hour_calib_shrink: float = 30.0     # shrinkage (days) for hour-of-day calibration
    large_quantile: float = 0.99
    large_min: float = 0.0
    round_unit: float = 1_000_000.0     # 0 disables the round-amount rule
    mc_sims: int = 1000
    mc_lookback_weeks: int = 8
    large_lookback_weeks: int = 26
    op_hour_threshold: float = 0.05
    objective: str = "wape"             # wape | balanced
    outlook_hours: int = 24
    seed: int = 42
    quick: bool = False


def parse_args() -> Config:
    ap = argparse.ArgumentParser(description="Chronos-2 + Monte Carlo 60-min forecaster")
    d = Config()
    for f, v in asdict(d).items():
        name = "--" + f.replace("_", "-")
        if isinstance(v, bool):
            ap.add_argument(name, action="store_true")
        else:
            ap.add_argument(name, type=type(v) if v is not None else str, default=v,
                            required=False)
    cfg = Config(**vars(ap.parse_args()))
    cfg.fm_context_hours = max(int(x) for x in cfg.fm_contexts.split(","))
    return cfg


def log(msg: str) -> None:
    print(f"[{time.strftime('%H:%M:%S')}] {msg}", flush=True)


# ============================================================================ loading
def _parse_bucket(b: pd.Series) -> pd.Series:
    """bucket may be 'HH:MM', 'HH:MM:SS', datetime.time, a Timestamp or an Excel day-fraction."""
    if pd.api.types.is_numeric_dtype(b):
        return pd.to_timedelta(((b.astype(float) % 1) * 1440).round(), unit="min")
    s = b.astype(str).str.extract(r"(\d{1,2}:\d{2})")[0]
    parts = s.str.split(":", expand=True)
    mins = pd.to_numeric(parts[0], errors="coerce") * 60 + pd.to_numeric(parts[1], errors="coerce")
    return pd.to_timedelta(mins, unit="min")


def load_minute_data(cfg: Config) -> tuple[pd.DataFrame, dict]:
    p = Path(cfg.input)
    if not p.exists():
        raise FileNotFoundError(f"Input file not found: {p}  (pass --input /full/path/to/file.xlsx)")
    if p.suffix.lower() in {".xlsx", ".xlsm", ".xls"}:
        sheet = cfg.sheet
        if not sheet:
            names = pd.ExcelFile(p).sheet_names
            sheet = "Sheet2" if "Sheet2" in names else names[0]
        log(f"Reading {p} (sheet '{sheet}') - large workbooks can take a few minutes ...")
        df = pd.read_excel(p, sheet_name=sheet)
    else:
        df = pd.read_csv(p)
    df.columns = [str(c).strip().lower() for c in df.columns]
    need = {"value_date", "bucket", "credit_amount", "debit_amount"}
    if need - set(df.columns):
        raise ValueError(f"Missing columns {need - set(df.columns)}; found {list(df.columns)}")
    dq = {"raw_rows": int(len(df))}
    ts = pd.to_datetime(df["value_date"], errors="coerce").dt.normalize() + _parse_bucket(df["bucket"])
    out = pd.DataFrame({"ts": ts,
                        "credit": pd.to_numeric(df["credit_amount"], errors="coerce"),
                        "debit": pd.to_numeric(df["debit_amount"], errors="coerce")})
    dq["unparseable_rows_dropped"] = int(out["ts"].isna().sum())
    out = out.dropna(subset=["ts"])
    dq["negative_values_made_absolute"] = int((out[["credit", "debit"]] < 0).sum().sum())
    out[["credit", "debit"]] = out[["credit", "debit"]].fillna(0.0).abs()
    dq["duplicate_minutes_merged"] = int(out["ts"].duplicated().sum())
    out = out.groupby("ts", as_index=False).sum().sort_values("ts").reset_index(drop=True)
    dq.update(minute_rows=int(len(out)), start=str(out.ts.min()), end=str(out.ts.max()),
              total_credit=float(out.credit.sum()), total_debit=float(out.debit.sum()))
    if out.ts.max().minute < 59:
        dq["warning_last_hour"] = "last hour may be partial (last minute < :59)"
    return out, dq



def load_schedule(cfg: Config) -> pd.DataFrame:
    """Optional known/scheduled payments: columns datetime, credit_amount, debit_amount."""
    if not cfg.schedule_file:
        return pd.DataFrame({"ts": pd.Series(dtype="datetime64[ns]"), "credit": pd.Series(dtype=float),
                             "debit": pd.Series(dtype=float)})
    s = pd.read_csv(cfg.schedule_file)
    s.columns = [c.strip().lower() for c in s.columns]
    ts = pd.to_datetime(s.iloc[:, 0], errors="coerce").dt.floor("min")
    out = pd.DataFrame({"ts": ts,
                        "credit": pd.to_numeric(s.get("credit_amount", 0), errors="coerce"),
                        "debit": pd.to_numeric(s.get("debit_amount", 0), errors="coerce")}).dropna(subset=["ts"])
    out[list(SERIES)] = out[list(SERIES)].astype(float).fillna(0).abs()
    return out.groupby("ts", as_index=False).sum()


def classify_minutes(m: pd.DataFrame, sched: pd.DataFrame, train_end, cfg: Config):
    """Split every minute amount into scheduled / large / routine. Thresholds from TRAIN only."""
    m = m.merge(sched.rename(columns={s: f"{s}_schedfile" for s in SERIES}), on="ts", how="left")
    thr = {}
    for s in SERIES:
        sf = pd.to_numeric(m[f"{s}_schedfile"], errors="coerce").fillna(0).to_numpy(dtype=float)
        amt = m[s].to_numpy(dtype=float)
        m[f"{s}_sched"] = np.minimum(sf, amt)                      # matched scheduled part
        rest = amt - m[f"{s}_sched"].to_numpy()
        tr = (m.ts < train_end).to_numpy() & (rest > 0)
        q = float(np.quantile(rest[tr], cfg.large_quantile)) if tr.any() else np.inf
        thr[s] = max(q, cfg.large_min)
        is_round = (cfg.round_unit > 0) & (rest >= cfg.round_unit) & (np.abs(rest / cfg.round_unit - np.round(rest / cfg.round_unit)) < 1e-9)
        is_large = (rest >= thr[s]) | is_round
        m[f"{s}_L"] = np.where(is_large, rest, 0.0)
        m[f"{s}_r"] = np.where(is_large, 0.0, rest)
        m[f"{s}_round"] = np.where(is_round, rest, 0.0)
    return m.drop(columns=[f"{s}_schedfile" for s in SERIES]), thr


def build_hourly(m: pd.DataFrame, sched: pd.DataFrame, end_extra_hours: int) -> tuple[pd.DataFrame, pd.Series]:
    m = m.assign(hour=m.ts.dt.floor("h"))
    late = m.ts.dt.minute >= 45
    agg = {}
    for s in SERIES:
        for c in (s, f"{s}_r", f"{s}_L", f"{s}_sched", f"{s}_round"):
            agg[c] = (c, "sum")
        m[f"{s}_n"] = (m[f"{s}_r"] > 0).astype(int)          # routine active minutes
        m[f"{s}_Ln"] = (m[f"{s}_L"] > 0).astype(int)         # large payments
        m[f"{s}_l15"] = np.where(late, m[f"{s}_r"], 0.0)
        agg[f"{s}_n"] = (f"{s}_n", "sum")
        agg[f"{s}_Ln"] = (f"{s}_Ln", "sum")
        agg[f"{s}_l15"] = (f"{s}_l15", "sum")
        agg[f"{s}_max"] = (f"{s}_r", "max")
    H = m.groupby("hour").agg(**agg)
    idx = pd.date_range(H.index.min().normalize(), H.index.max(), freq="h")
    H = H.reindex(idx, fill_value=0.0)
    H.index.name = "hour"
    # scheduled amounts known in advance for FUTURE hours (outlook / next hour)
    fut = {}
    if len(sched):
        sh = sched.assign(hour=sched.ts.dt.floor("h")).groupby("hour")[list(SERIES)].sum()
        fut = sh
    return H, fut


# ============================================================================ calendar & metrics
def load_holidays(cfg: Config, years) -> set:
    hol = set()
    if cfg.holiday_country:
        if pyholidays is None:
            log("WARNING: `holidays` package not installed -> country holidays skipped "
                "(pip install holidays, or pass --holiday-file)")
        else:
            hol |= {pd.Timestamp(d) for d in pyholidays.country_holidays(cfg.holiday_country, years=list(years))}
    if cfg.holiday_file:
        hf = pd.read_csv(cfg.holiday_file)
        hol |= set(pd.to_datetime(hf.iloc[:, 0], errors="coerce").dropna().dt.normalize())
    return hol


def calendar_features(idx: pd.DatetimeIndex, hol: set) -> pd.DataFrame:
    """Deterministic calendar features - known in advance, so leakage-free."""
    dates = idx.normalize()
    full = pd.date_range(dates.min() - pd.offsets.MonthBegin(1), dates.max() + pd.offsets.MonthEnd(1)
                         + pd.Timedelta(days=15), freq="D")
    D = pd.DataFrame(index=full)
    D["is_weekend"] = full.dayofweek >= 5
    D["is_holiday"] = full.isin(list(hol))
    D["is_bday"] = ~(D.is_weekend | D.is_holiday)
    ym = full.to_period("M")
    bd_cum = D.groupby(ym)["is_bday"].cumsum()
    bd_tot = D.groupby(ym)["is_bday"].transform("sum")
    D["bd_of_month"] = np.where(D.is_bday, bd_cum, 0)
    D["bd_to_month_end"] = np.where(D.is_bday, bd_tot - bd_cum, -1)
    D["is_month_start_bd"] = D.bd_of_month == 1
    D["is_month_end_bd"] = D.is_bday & (D.bd_to_month_end == 0)
    D["is_quarter_end_bd"] = D.is_month_end_bd & full.month.isin([3, 6, 9, 12])
    D["is_year_end_bd"] = D.is_month_end_bd & (full.month == 12)
    D["is_mid_month"] = full.day.isin([14, 15, 16])
    D["next_day_closed"] = (~D.is_bday).shift(-1, fill_value=False)   # Friday / pre-holiday
    D["prev_day_closed"] = (~D.is_bday).shift(1, fill_value=False)    # Monday / post-holiday
    hd = full[D.is_holiday.values].values
    if len(hd):
        pos = np.searchsorted(hd, full.values)
        nxt = np.where(pos < len(hd), (hd[np.minimum(pos, len(hd) - 1)] - full.values) / np.timedelta64(1, "D"), 99)
        prv = np.where(pos > 0, (full.values - hd[np.maximum(pos - 1, 0)]) / np.timedelta64(1, "D"), 99)
        D["days_to_holiday"] = np.clip(nxt, 0, 15)
        D["days_since_holiday"] = np.clip(prv, 0, 15)
    else:
        D["days_to_holiday"] = 15
        D["days_since_holiday"] = 15
    C = D.reindex(dates)
    C.index = idx
    C["hour"] = idx.hour
    C["dow"] = idx.dayofweek
    C["dom"] = idx.day
    C["month"] = idx.month
    C["quarter"] = idx.quarter
    C["week_of_year"] = idx.isocalendar().week.values.astype(int)
    C["how"] = C.dow * 24 + C.hour                                     # hour-of-week
    for col, per in (("hour", 24), ("dow", 7), ("dom", 31), ("month", 12)):
        C[f"{col}_sin"] = np.sin(2 * np.pi * C[col] / per)
        C[f"{col}_cos"] = np.cos(2 * np.pi * C[col] / per)
    return C.astype(float)


# ----------------------------------------------------------------------------- metrics
def metrics(a, f) -> dict:
    a, f = np.asarray(a, float), np.asarray(f, float)
    e = f - a
    den = np.abs(a).sum()
    wape = np.abs(e).sum() / den if den > 0 else np.nan
    s_den = np.abs(a) + np.abs(f)
    smape = np.mean(np.where(s_den > 0, 2 * np.abs(e) / np.where(s_den > 0, s_den, 1), 0))
    return {"Accuracy_%": 100 * (1 - wape), "WAPE_%": 100 * wape, "MAE": np.abs(e).mean(),
            "RMSE": np.sqrt((e ** 2).mean()), "Bias_%": 100 * e.sum() / den if den > 0 else np.nan,
            "sMAPE_%": 100 * smape}


def wape(a, f) -> float:
    a, f = np.asarray(a, float), np.asarray(f, float)
    d = np.abs(a).sum()
    return np.abs(f - a).sum() / d if d > 0 else np.inf


def operating_hours(H: pd.DataFrame, mask, thr: float) -> list[int]:
    active = (H["credit_n"] + H["debit_n"] + H["credit_Ln"] + H["debit_Ln"]) > 0
    frac = active[mask].groupby(H.index[mask].hour).mean()
    return sorted(int(h) for h in frac[frac >= thr].index)



# ============================================================================ features (GBM)
def build_features(H: pd.DataFrame, cal: pd.DataFrame, sched_h: pd.DataFrame, cfg: Config) -> pd.DataFrame:
    """Features for hour h use observations <= h-k (k = 1 + latency). Calendar and the
    scheduled-payment file are known in advance, so their value AT h is allowed."""
    k = 1 + cfg.latency_hours
    X = pd.DataFrame(index=H.index)
    day = pd.Series(H.index.normalize(), index=H.index)
    same_day = day.shift(k) == day
    for s in SERIES:
        y = np.log1p(H[f"{s}_r"])
        p = y.shift(k)
        for L in [L for L in (1, 2, 3, 4, 24, 48, 168, 336) if L >= k]:
            X[f"{s}_lag{L}"] = y.shift(L)
        for w in (3, 6, 12, 24, 72, 168):
            X[f"{s}_rmean{w}"] = p.rolling(w, min_periods=1).mean()
        X[f"{s}_rstd24"] = p.rolling(24, min_periods=2).std()
        X[f"{s}_ewm6"] = p.ewm(halflife=6, min_periods=1).mean()
        wk = pd.concat([y.shift(168 * i) for i in range(1, 5)], axis=1)
        X[f"{s}_sh4w_mean"] = wk.mean(axis=1)
        X[f"{s}_sh4w_med"] = wk.median(axis=1)
        X[f"{s}_sh4w_std"] = wk.std(axis=1)
        X[f"{s}_sh5d_mean"] = pd.concat([y.shift(24 * i) for i in range(1, 6)], axis=1).mean(axis=1)
        n = H[f"{s}_n"]
        X[f"{s}_n_lag"] = n.shift(k)
        X[f"{s}_n_rmean24"] = n.shift(k).rolling(24, min_periods=1).mean()
        X[f"{s}_n_sh4w"] = pd.concat([n.shift(168 * i) for i in range(1, 5)], axis=1).mean(axis=1)
        X[f"{s}_last15_lag"] = np.log1p(H[f"{s}_l15"].shift(k))
        X[f"{s}_max_lag"] = np.log1p(H[f"{s}_max"].shift(k))
        cs = H[f"{s}_r"].groupby(day).cumsum()
        sofar = np.log1p(cs.shift(k).where(same_day, 0.0))
        X[f"{s}_today_sofar"] = sofar
        X[f"{s}_today_vs_lastweek"] = sofar - sofar.shift(168)
        # lumpy-payment context
        X[f"{s}_large_lag"] = np.log1p(H[f"{s}_L"].shift(k))
        X[f"{s}_large_cnt24"] = H[f"{s}_Ln"].shift(k).rolling(24, min_periods=1).sum()
        X[f"{s}_large_cnt168"] = H[f"{s}_Ln"].shift(k).rolling(168, min_periods=1).sum()
        X[f"{s}_round_24"] = np.log1p(H[f"{s}_round"].shift(k).rolling(24, min_periods=1).sum())
        X[f"{s}_total_lag"] = np.log1p(H[s].shift(k))
        X[f"{s}_sched_now"] = np.log1p(sched_h[s].reindex(H.index).fillna(0))   # known in advance
    net = (H["credit"] - H["debit"]).shift(k)
    X["net_lag"] = np.sign(net) * np.log1p(net.abs())
    return X.join(cal)


def fm_windows(y: np.ndarray, positions, C: int, lat: int):
    """Context window for target position p = y[p-lat-C : p-lat]  (strictly before origin)."""
    pos = np.asarray(positions)
    start = pos - lat - C
    idx = start[:, None] + np.arange(C)[None, :]
    return y[np.clip(idx, 0, None)] * (idx >= 0)


def leakage_check(H, cal, sched_h, cfg, n=4) -> list[dict]:
    rng = np.random.default_rng(cfg.seed)
    base = build_features(H, cal, sched_h, cfg)
    res = []
    for pos in rng.integers(cfg.warmup_hours, len(H) - 2, size=n):
        t = H.index[pos]
        H2 = H.copy()
        fut = H2.index >= t
        H2.loc[fut] = rng.random((int(fut.sum()), H2.shape[1])) * 1e9
        X2 = build_features(H2, cal, sched_h, cfg)
        rows = base.index <= t
        ok_f = bool(np.allclose(np.nan_to_num(base.loc[rows].to_numpy(), nan=-7.7),
                                np.nan_to_num(X2.loc[rows].to_numpy(), nan=-7.7)))
        y1, y2 = np.log1p(H["credit_r"].to_numpy()), np.log1p(H2["credit_r"].to_numpy())
        ok_w = bool(np.array_equal(fm_windows(y1, [pos], cfg.fm_context_hours, cfg.latency_hours),
                                   fm_windows(y2, [pos], cfg.fm_context_hours, cfg.latency_hours)))
        res.append({"cut_hour": str(t), "gbm_features_unchanged": ok_f, "fm_context_unchanged": ok_w})
    return res


# ============================================================================ gradient boosting
def gbm_engine(cfg: Config) -> str:
    if cfg.gbm_backend != "auto":
        return cfg.gbm_backend
    return "lightgbm" if lgb else ("xgboost" if xgb else "sklearn")


def param_grid(quick: bool) -> list[dict]:
    g = [dict(n=500, lr=0.04, leaves=31, min_leaf=40, l2=1.0),
         dict(n=900, lr=0.03, leaves=63, min_leaf=25, l2=3.0),
         dict(n=400, lr=0.05, leaves=15, min_leaf=80, l2=1.0)]
    return g[:1] if quick else g


def make_model(p: dict, cfg: Config):
    eng = gbm_engine(cfg)
    if eng == "lightgbm":
        return lgb.LGBMRegressor(n_estimators=p["n"], learning_rate=p["lr"], num_leaves=p["leaves"],
                                 min_child_samples=p["min_leaf"], reg_lambda=p["l2"], subsample=0.8,
                                 subsample_freq=1, colsample_bytree=0.8, random_state=cfg.seed, verbose=-1)
    if eng == "xgboost":
        return xgb.XGBRegressor(n_estimators=p["n"], learning_rate=p["lr"], max_leaves=p["leaves"],
                                grow_policy="lossguide", tree_method="hist", min_child_weight=p["min_leaf"] / 10,
                                reg_lambda=p["l2"], subsample=0.8, colsample_bytree=0.8, random_state=cfg.seed)
    return HistGradientBoostingRegressor(max_iter=p["n"], learning_rate=p["lr"], max_leaf_nodes=p["leaves"],
                                         min_samples_leaf=p["min_leaf"], l2_regularization=p["l2"],
                                         early_stopping=False, random_state=cfg.seed)


def walk_forward(X, y, cols, fit_start, seg_start, seg_end, p, cfg):
    idx = X.index
    out = pd.Series(np.nan, index=idx[(idx >= seg_start) & (idx <= seg_end)])
    cur, model = seg_start, None
    purge = pd.Timedelta(hours=cfg.latency_hours)
    while cur <= seg_end:
        nxt = min(cur + pd.Timedelta(days=cfg.refit_every_days), seg_end + HOUR)
        tr = (idx >= fit_start) & (idx < cur - purge) & y.notna().to_numpy()
        model = make_model(p, cfg).fit(X.loc[tr, cols], y[tr])
        blk = (idx >= cur) & (idx < nxt)
        out.loc[idx[blk]] = model.predict(X.loc[blk, cols])
        cur = nxt
    return out.clip(lower=0), model


# ============================================================================ Chronos-2
FM_COVS = ("hour_sin", "hour_cos", "dow_sin", "dow_cos", "is_weekend", "is_holiday", "is_month_end_bd",
           "is_month_start_bd", "is_quarter_end_bd", "next_day_closed", "prev_day_closed", "bd_to_month_end")
PAST_COVS = ("credit_large", "debit_large")          # history-only covariates (unknown in future)
FM_STATS = ("mean", "q10", "q20", "q30", "q40", "q50", "q60", "q70", "q80", "q90")
Q_GRID = [0.1, 0.2, 0.3, 0.4, 0.5, 0.6, 0.7, 0.8, 0.9]


class Chronos2Forecaster:
    """Amazon Chronos-2 (zero-shot, covariate-aware) via Chronos2Pipeline.predict_df.
       * target  : credit & debit routine flow (log1p) forecast JOINTLY (multivariate)
       * future-known covariates : calendar features (context + forecast hours)
       * past-only covariates    : large-payment history (log1p)
       Each forecast origin is sent as its own series whose history ends strictly
       before the origin, so no future value can enter the context."""

    def __init__(self, cfg: Config):
        import chronos
        from chronos import Chronos2Pipeline
        self.cfg = cfg
        self.is_mock = bool(getattr(chronos, "__mock__", False))
        dev = cfg.fm_device
        if dev == "auto":
            try:
                import torch
                dev = "cuda" if torch.cuda.is_available() else "cpu"
            except ImportError:
                dev = "cpu"
        self.device = dev
        self.pipe = Chronos2Pipeline.from_pretrained(cfg.fm_model, device_map=dev)
        self.cov_status = "calendar (future-known) + large-payment history (past-only), joint credit/debit"

    @staticmethod
    def _parse(pred: pd.DataFrame, n: int, steps: int) -> dict:
        """pred_df -> {series: (n, steps, 10)} with [mean, q10..q90]."""
        cols = {str(c): c for c in pred.columns}
        qcol = {}
        for q in Q_GRID:
            for key in (str(q), f"{q:.1f}", f"q{q}", f"{q}"):
                if key in cols:
                    qcol[q] = cols[key]
                    break
        pcol = next((cols[k] for k in ("predictions", "mean", "prediction") if k in cols), qcol.get(0.5))
        tcol = next((cols[k] for k in ("target_name", "target", "variable") if k in cols
                     and not pd.api.types.is_numeric_dtype(pred[cols[k]])), None)
        if tcol is None and len(pred) == 2 * n * steps:
            raise ValueError(f"Cannot identify the target column in Chronos-2 output: {list(pred.columns)}")
        out = {}
        for s in SERIES:
            sub = pred[pred[tcol].astype(str) == s] if tcol is not None else pred
            sub = sub.sort_values(["id", "timestamp"])
            arr = np.column_stack([sub[pcol].to_numpy(float)] + [sub[qcol[q]].to_numpy(float) if q in qcol
                                                                  else sub[pcol].to_numpy(float) for q in Q_GRID])
            out[s] = arr.reshape(n, steps, 10)
        return out

    def predict(self, H, cal, positions, C: int, horizon: int) -> dict:
        lat = self.cfg.latency_hours
        steps = lat + horizon
        y = {s: np.log1p(H[f"{s}_r"].to_numpy()) for s in SERIES}
        lg = {f"{s}_large": np.log1p(H[f"{s}_L"].to_numpy()) for s in SERIES}
        calm = cal[list(FM_COVS)].to_numpy(float)
        cal_ts = cal.index.values
        res = {s: [] for s in SERIES}
        pos_all = np.asarray(positions)
        for i in range(0, len(pos_all), self.cfg.fm_batch):
            pos = pos_all[i:i + self.cfg.fm_batch]
            Ceff = int(min(C, pos.min() - lat))
            ci = (pos - lat - Ceff)[:, None] + np.arange(Ceff)[None, :]          # strictly < origin - lat
            ids = np.repeat(np.arange(len(pos)), Ceff)
            ctx = pd.DataFrame({"id": ids, "timestamp": cal_ts[ci.ravel()]})
            for s in SERIES:
                ctx[s] = y[s][ci.ravel()]
            for k, v in lg.items():
                ctx[k] = v[ci.ravel()]
            for j, c in enumerate(FM_COVS):
                ctx[c] = calm[ci.ravel(), j]
            fi = (pos - lat)[:, None] + np.arange(steps)[None, :]
            fut = pd.DataFrame({"id": np.repeat(np.arange(len(pos)), steps), "timestamp": cal_ts[fi.ravel()]})
            for j, c in enumerate(FM_COVS):
                fut[c] = calm[fi.ravel(), j]
            pred = self.pipe.predict_df(ctx, future_df=fut, prediction_length=steps, quantile_levels=Q_GRID,
                                        id_column="id", timestamp_column="timestamp", target=list(SERIES))
            parsed = self._parse(pred, len(pos), steps)
            for s in SERIES:
                res[s].append(parsed[s][:, lat:, :])
        return {s: np.clip(np.concatenate(v), 0, None) for s, v in res.items()}


def load_fm(cfg: Config):
    if cfg.fm_backend == "none":
        return None, "disabled (--fm-backend none)"
    try:
        fm = Chronos2Forecaster(cfg)
        return fm, ("MOCK stand-in for Chronos-2 (sandbox test only - NOT the real model)" if fm.is_mock
                    else f"Chronos-2 ({cfg.fm_model}, {fm.device})")
    except Exception as e:  # noqa: BLE001
        if cfg.fm_backend == "chronos":
            raise RuntimeError('Chronos-2 is required: pip install "chronos-forecasting>=2.0" "pandas[pyarrow]" '
                               "(or run with --fm-backend auto to continue without it)") from e
        log(f"WARNING Chronos-2 unavailable ({type(e).__name__}: {str(e)[:100]}) -> continuing without it")
        return None, f"not available ({type(e).__name__})"


def build_pools(minute: pd.DataFrame, suffix: str = "_r") -> dict:
    m = minute.assign(hour=minute.ts.dt.floor("h"))
    pools = {}
    for s in SERIES:
        c = f"{s}{suffix}"
        sub = m.loc[m[c] > 0, ["hour", c]]
        pools[s] = {h: g.to_numpy() for h, g in sub.groupby("hour")[c]}
    return pools


class MonteCarloStructural:
    """ROUTINE flow. Compound-Poisson simulation per target hour:
         N ~ Poisson(lambda * calendar_factor)   lambda = mean active minutes, same hour-of-week, last K normal weeks
         amounts ~ bootstrap of those weeks' minute amounts * calendar amount factor
       Calendar factors estimated only on data < fit_end."""
    FLAGS = ("is_holiday", "is_quarter_end_bd", "is_month_end_bd", "is_month_start_bd")

    def __init__(self, pools, H, cal, fit_end, cfg: Config):
        self.pools, self.H, self.cal, self.cfg = pools, H, cal, cfg
        self.rng = np.random.default_rng(cfg.seed)
        self.special = cal[list(self.FLAGS)].loc[H.index].any(axis=1).to_numpy()
        how = cal["how"].loc[H.index].to_numpy().astype(int)
        normal_idx = H.index[~self.special]
        self.hist = {w: normal_idx[how[~self.special] == w].values for w in range(168)}
        self.factors = self._factors(fit_end)

    def _factors(self, fit_end):
        H, cal = self.H, self.cal.loc[self.H.index]
        fit = (H.index < fit_end)
        f = {}
        for s in SERIES:
            n, amt = H[f"{s}_n"], H[f"{s}_r"]
            normal = fit & ~self.special
            exp_n = n[normal].groupby(cal["how"][normal]).mean()
            exp_a = amt[normal].groupby(cal["how"][normal]).sum() / n[normal].groupby(cal["how"][normal]).sum().replace(0, np.nan)
            f[s] = {}
            for flag in self.FLAGS:
                rows = fit & (cal[flag].to_numpy() > 0)
                hw = cal["how"][rows]
                en = exp_n.reindex(hw).to_numpy()
                ea = exp_a.reindex(hw).to_numpy()
                n_act, a_act = n[rows].to_numpy(), amt[rows].to_numpy()
                if np.nansum(en) > 20 and n_act.sum() > 20:
                    cf = n_act.sum() / np.nansum(en)
                    mean_amt_flag = a_act.sum() / max(n_act.sum(), 1)
                    mean_amt_norm = np.nansum(ea * n_act) / max(n_act.sum(), 1)
                    af = mean_amt_flag / mean_amt_norm if mean_amt_norm > 0 else 1.0
                else:
                    cf, af = 1.0, 1.0
                f[s][flag] = (float(np.clip(cf, 0.02, 5)), float(np.clip(af, 0.1, 10)))
        return f

    def _factor_for(self, s, h):
        c = self.cal.loc[h]
        if c["is_holiday"] > 0:
            return self.factors[s]["is_holiday"]
        cf, af = 1.0, 1.0
        for flag in ("is_quarter_end_bd", "is_month_end_bd", "is_month_start_bd"):
            if c[flag] > 0:
                cf, af = cf * self.factors[s][flag][0], af * self.factors[s][flag][1]
                if flag == "is_quarter_end_bd":
                    break               # quarter-end already contains the month-end effect
        return cf, af

    def forecast(self, hours, cutoff=None) -> dict:
        K, S = self.cfg.mc_lookback_weeks, self.cfg.mc_sims
        lat = pd.Timedelta(hours=self.cfg.latency_hours)
        res = {s: {q: np.zeros(len(hours)) for q in ("p10", "p50", "p90", "mean")} for s in SERIES}
        pos_of = pd.Series(np.arange(len(self.H)), index=self.H.index)
        rep = np.arange(S)
        for i, h in enumerate(hours):
            c = cutoff if cutoff is not None else h - lat
            how = int(self.cal.at[h, "how"])
            cand = self.hist[how]
            j = np.searchsorted(cand, np.datetime64(c))           # strictly before the origin
            sel = cand[max(0, j - K):j]
            assert len(sel) == 0 or sel.max() < np.datetime64(c)
            if len(sel) == 0:
                continue
            for s in SERIES:
                lam = self.H[f"{s}_n"].to_numpy()[pos_of[sel].to_numpy()].mean()
                cf, af = self._factor_for(s, h)
                pool = [self.pools[s][t] for t in pd.DatetimeIndex(sel) if t in self.pools[s]]
                if lam <= 0 or not pool:
                    continue
                pool = np.concatenate(pool) * af
                N = self.rng.poisson(lam * cf, S)
                draws = pool[self.rng.integers(0, len(pool), N.sum())]
                sims = np.bincount(np.repeat(rep, N), weights=draws, minlength=S)
                q = np.percentile(sims, [10, 50, 90])
                r = res[s]
                r["p10"][i], r["p50"][i], r["p90"][i], r["mean"][i] = q[0], q[1], q[2], sims.mean()
        return {s: {k: pd.Series(v, index=hours) for k, v in d.items()} for s, d in res.items()}


# ----------------------------------------------------------------------------- hybrid blend
def simplex(n: int, step: float = 0.05):
    k = int(round(1 / step))
    for c in itertools.product(range(k + 1), repeat=n - 1):
        if sum(c) <= k:
            yield np.array(list(c) + [k - sum(c)]) / k


def fit_blend(actual, comps: dict, mask) -> dict:
    names = list(comps)
    M = np.column_stack([np.asarray(comps[n], float)[mask] for n in names])
    a = np.asarray(actual, float)[mask]
    best_w, best = None, np.inf
    for w in simplex(len(names)):
        e = wape(a, M @ w)
        if e < best:
            best, best_w = e, w
    return {n: float(w) for n, w in zip(names, best_w)}


def fit_scale(actual, fc, mask, objective: str, max_bias: float = 0.05) -> float:
    """Multiplicative calibration learned on VALIDATION. Log-target models forecast the median,
    which under-states totals of heavy-tailed flows; 'balanced' removes that bias."""
    a, f = np.asarray(actual, float)[mask], np.asarray(fc, float)[mask]
    grid = np.round(np.arange(0.5, 4.001, 0.01), 2)
    stats = [(c, wape(a, c * f), (c * f - a).sum() / max(np.abs(a).sum(), 1e-9)) for c in grid]
    if objective == "wape":
        return float(min(stats, key=lambda t: t[1])[0])
    ok = [t for t in stats if abs(t[2]) <= max_bias]
    return float(min(ok, key=lambda t: t[1])[0] if ok else min(stats, key=lambda t: abs(t[2]))[0])


def apply_blend(comps: dict, w: dict) -> np.ndarray:
    return sum(w[n] * np.asarray(comps[n], float) for n in w)


class ResidualMonteCarlo:
    """Joint bootstrap of validation log-residuals (credit & debit sampled together, so their
    correlation carries into NET). Residuals are grouped by hour-of-day."""

    def __init__(self, val_idx, actual: dict, fc: dict, cfg: Config, min_group=30):
        self.rng = np.random.default_rng(cfg.seed + 1)
        self.S = cfg.mc_sims
        self.R = np.column_stack([np.log1p(np.asarray(actual[s], float)) - np.log1p(np.asarray(fc[s], float))
                                  for s in SERIES])
        hod = np.asarray(val_idx.hour)
        self.groups = {h: np.where(hod == h)[0] for h in range(24)}
        self.pooled = np.arange(len(self.R))
        self.min_group = min_group

    def simulate(self, idx, fc: dict) -> dict:
        n = len(idx)
        out = {k: {q: np.zeros(n) for q in ("p10", "p50", "p90")} for k in (*SERIES, "net")}
        hod = np.asarray(idx.hour)
        for i in range(n):
            g = self.groups.get(hod[i])
            g = g if g is not None and len(g) >= self.min_group else self.pooled
            r = self.R[self.rng.choice(g, self.S)]
            sims = {}
            for j, s in enumerate(SERIES):
                sims[s] = np.clip(np.expm1(np.log1p(max(fc[s][i], 0)) + r[:, j]), 0, None)
            sims["net"] = sims["credit"] - sims["debit"]
            for k, v in sims.items():
                q = np.percentile(v, [10, 50, 90])
                out[k]["p10"][i], out[k]["p50"][i], out[k]["p90"][i] = q
        return {k: {q: pd.Series(v, index=idx) for q, v in d.items()} for k, d in out.items()}




class LargePaymentMC:
    """LARGE payments: probability x size Monte Carlo.
       rate  = mean large-payment count, same hour-of-day and day type (business / closed),
               last `large_lookback_weeks` weeks strictly before the origin, x calendar factor
       sizes = bootstrap of large payments in the same look-back window, x calendar size factor
       Calendar factors estimated on data < fit_end only."""
    FLAGS = ("is_quarter_end_bd", "is_month_end_bd", "is_month_start_bd", "prev_day_closed")

    def __init__(self, minute, H, cal, fit_end, cfg: Config):
        self.H, self.cal, self.cfg = H, cal, cfg
        self.rng = np.random.default_rng(cfg.seed + 7)
        self.S = cfg.mc_sims
        self.W = pd.Timedelta(weeks=cfg.large_lookback_weeks)
        c = cal.loc[H.index]
        self.key = (c["hour"].astype(int) * 2 + c["is_bday"].astype(int)).to_numpy()
        self.hours_by_key = {k: H.index[self.key == k].values for k in np.unique(self.key)}
        self.pos = pd.Series(np.arange(len(H)), index=H.index)
        self.sizes = {}
        for s in SERIES:
            m = minute.loc[minute[f"{s}_L"] > 0, ["ts", f"{s}_L"]].sort_values("ts")
            self.sizes[s] = (m.ts.values, m[f"{s}_L"].to_numpy())
        self.factors = self._factors(fit_end)

    def _factors(self, fit_end):
        H, c = self.H, self.cal.loc[self.H.index]
        fit = (H.index < fit_end) & (c["is_bday"] > 0).to_numpy()
        out = {}
        for s in SERIES:
            n, a = H[f"{s}_Ln"], H[f"{s}_L"]
            flagged = c[list(self.FLAGS)].any(axis=1).to_numpy()
            normal = fit & ~flagged
            base_rate = n[normal].groupby(c["hour"][normal]).mean()
            base_size = a[normal].sum() / max(n[normal].sum(), 1)
            out[s] = {}
            for f in self.FLAGS:
                rows = fit & (c[f] > 0).to_numpy()
                exp = base_rate.reindex(c["hour"][rows]).to_numpy()
                if np.nansum(exp) > 5 and n[rows].sum() > 5:
                    rf = n[rows].sum() / np.nansum(exp)
                    sf = (a[rows].sum() / max(n[rows].sum(), 1)) / base_size if base_size > 0 else 1
                else:
                    rf, sf = 1.0, 1.0
                out[s][f] = (float(np.clip(rf, 0.05, 10)), float(np.clip(sf, 0.2, 5)))
        return out

    def forecast(self, hours, cutoff=None) -> dict:
        res = {s: {q: np.zeros(len(hours)) for q in ("mean", "p10", "p50", "p90", "prob")} for s in SERIES}
        lat = pd.Timedelta(hours=self.cfg.latency_hours)
        rep = np.arange(self.S)
        for i, h in enumerate(hours):
            co = np.datetime64(cutoff if cutoff is not None else h - lat)
            ch = self.cal.loc[h]
            key = int(ch["hour"]) * 2 + int(ch["is_bday"])
            cand = self.hours_by_key.get(key, np.array([], dtype="datetime64[ns]"))
            j = np.searchsorted(cand, co)
            sel = cand[(cand < co) & (cand >= co - np.timedelta64(self.W))]
            if len(sel) == 0:
                continue
            for s in SERIES:
                rate = self.H[f"{s}_Ln"].to_numpy()[self.pos[sel].to_numpy()].mean()
                rf, sf = 1.0, 1.0
                for f in self.FLAGS:
                    if ch[f] > 0:
                        rf, sf = rf * self.factors[s][f][0], sf * self.factors[s][f][1]
                ts, amt = self.sizes[s]
                w = (ts < co) & (ts >= co - np.timedelta64(self.W))
                pool = amt[w]
                if rate <= 0 or len(pool) == 0:
                    continue
                N = self.rng.poisson(rate * rf, self.S)
                draws = pool[self.rng.integers(0, len(pool), N.sum())] * sf
                sims = np.bincount(np.repeat(rep, N), weights=draws, minlength=self.S)
                r = res[s]
                r["mean"][i] = sims.mean()
                r["p10"][i], r["p50"][i], r["p90"][i] = np.percentile(sims, Q_LEVELS)
                r["prob"][i] = 1 - np.exp(-rate * rf)
        return {s: {k: pd.Series(v, index=hours) for k, v in d.items()} for s, d in res.items()}


def fit_hour_scale(actual, fc, hours, mask, shrink_days: float) -> dict:
    """Per hour-of-day multiplicative correction, shrunk towards 1 when data are few."""
    a, f, h = np.asarray(actual, float), np.asarray(fc, float), np.asarray(hours)
    out = {}
    grid = np.round(np.arange(0.5, 2.001, 0.02), 2)
    for hr in np.unique(h[mask]):
        r = mask & (h == hr)
        if f[r].sum() <= 0 or a[r].sum() <= 0:
            continue
        best = min(grid, key=lambda c: wape(a[r], c * f[r]))
        w = r.sum() / (r.sum() + shrink_days)
        out[int(hr)] = float(1 + (best - 1) * w)
    return out


def apply_hour_scale(fc, hours, hs: dict) -> np.ndarray:
    return np.asarray(fc, float) * np.array([hs.get(int(h), 1.0) for h in hours])


def fit_large_scale(actual_total, routine_fc, large_mean, sched, mask, objective):
    a = np.asarray(actual_total, float)[mask]
    r, L, sc = (np.asarray(x, float)[mask] for x in (routine_fc, large_mean, sched))
    best = None
    for c in np.round(np.arange(0, 2.001, 0.05), 2):
        f = r + c * L + sc
        e, b = wape(a, f), (f - a).sum() / max(np.abs(a).sum(), 1e-9)
        score = e if objective == "wape" else (e if abs(b) <= 0.05 else 10 + abs(b))
        if best is None or score < best[0]:
            best = (score, c)
    return float(best[1])
# ----------------------------------------------------------------------------- charts
def fig_to_b64(fig) -> str:
    buf = io.BytesIO()
    fig.savefig(buf, format="png", dpi=110, bbox_inches="tight")
    plt.close(fig)
    return base64.b64encode(buf.getvalue()).decode()


def chart_hourly(res: pd.DataFrame, days: int = 7) -> str:
    r = res[res.index >= res.index.max() - pd.Timedelta(days=days)]
    fig, axes = plt.subplots(3, 1, figsize=(12, 9), sharex=True)
    for ax, k, col in zip(axes, ("credit", "debit", "net"), ("#1b6ca8", "#c0392b", "#2d8a4e")):
        ax.fill_between(r.index, r[f"{k}_p10"], r[f"{k}_p90"], color=col, alpha=.15, label="P10-P90 (Monte Carlo)")
        ax.plot(r.index, r[f"{k}_actual"], color="#222", lw=1.3, label="Actual")
        ax.plot(r.index, r[f"{k}_forecast"], color=col, lw=1.3, ls="--", label="Hybrid forecast")
        ax.set_title(f"{k.upper()} - hourly actual vs 60-min-ahead forecast (last {days} test days)", fontsize=10)
        ax.grid(alpha=.3)
        ax.legend(fontsize=8, loc="upper left")
    return fig_to_b64(fig)


def chart_daily(daily: pd.DataFrame) -> str:
    fig, axes = plt.subplots(2, 1, figsize=(12, 6), sharex=True)
    for ax, k in zip(axes, SERIES):
        ax.bar(daily.index, daily[f"{k}_actual"], color="#bbb", label="Actual (daily sum)")
        ax.plot(daily.index, daily[f"{k}_forecast"], color="#1b6ca8", marker=".", label="Sum of hourly forecasts")
        ax.set_title(f"{k.upper()} - daily totals over TEST period", fontsize=10)
        ax.legend(fontsize=8)
        ax.grid(alpha=.3)
    return fig_to_b64(fig)


def chart_by_hour(res: pd.DataFrame, op: list, snaive: dict) -> str:
    fig, axes = plt.subplots(1, 2, figsize=(12, 3.8))
    for ax, s in zip(axes, SERIES):
        r = res[res.index.hour.isin(op)]
        hyb = r.groupby(r.index.hour).apply(lambda g: 100 * (1 - wape(g[f"{s}_actual"], g[f"{s}_forecast"])))
        sn = snaive[s].loc[r.index]
        base = pd.DataFrame({"a": r[f"{s}_actual"], "f": sn}).groupby(r.index.hour).apply(
            lambda g: 100 * (1 - wape(g["a"], g["f"])))
        x = np.arange(len(hyb))
        ax.bar(x - .2, hyb.values, .4, label="Hybrid", color="#1b6ca8")
        ax.bar(x + .2, base.values, .4, label="Seasonal naive (t-1 week)", color="#bbb")
        ax.set_xticks(x, [f"{h:02d}" for h in hyb.index])
        ax.set_title(f"{s.upper()} accuracy (1-WAPE) by hour of day - TEST", fontsize=10)
        ax.set_ylabel("%")
        ax.grid(alpha=.3, axis="y")
        ax.legend(fontsize=8)
    return fig_to_b64(fig)


def chart_importance(imp: pd.Series) -> str:
    imp = imp.sort_values().tail(20)
    fig, ax = plt.subplots(figsize=(8, 6))
    ax.barh(imp.index, imp.values, color="#1b6ca8")
    ax.set_title("Top 20 features - CREDIT model", fontsize=10)
    ax.grid(alpha=.3, axis="x")
    return fig_to_b64(fig)


def chart_scatter(res: pd.DataFrame, op) -> str:
    r = res[res.index.hour.isin(op)]
    fig, axes = plt.subplots(1, 2, figsize=(11, 4.5))
    for ax, s in zip(axes, SERIES):
        a, f = np.log10(1 + r[f"{s}_actual"]), np.log10(1 + r[f"{s}_forecast"])
        ax.scatter(a, f, s=6, alpha=.35, color="#1b6ca8")
        lim = [0, max(a.max(), f.max()) * 1.02]
        ax.plot(lim, lim, color="#c0392b", lw=1)
        ax.set_xlabel("log10(actual)")
        ax.set_ylabel("log10(forecast)")
        ax.set_title(f"{s.upper()} - forecast vs actual (TEST, operating hours)", fontsize=10)
        ax.grid(alpha=.3)
    return fig_to_b64(fig)


def chart_routine(res: pd.DataFrame, days: int = 7) -> str:
    r = res[res.index >= res.index.max() - pd.Timedelta(days=days)]
    fig, axes = plt.subplots(2, 1, figsize=(12, 6), sharex=True)
    for ax, s, col in zip(axes, SERIES, ("#1b6ca8", "#c0392b")):
        ax.plot(r.index, r[f"{s}_routine_actual"], color="#222", lw=1.2, label="Routine actual")
        ax.plot(r.index, r[f"{s}_routine_forecast"], color=col, lw=1.2, ls="--", label="Routine forecast")
        big = r[r[f"{s}_large_actual"] > 0]
        ax.scatter(big.index, np.zeros(len(big)), marker="^", color="#f0b400", s=30, label="Large payment occurred")
        ax.set_title(f"{s.upper()} - ROUTINE flow (large payments removed), last {days} test days", fontsize=10)
        ax.legend(fontsize=8, loc="upper left")
        ax.grid(alpha=.3)
    return fig_to_b64(fig)


def chart_attribution(att: pd.DataFrame) -> str:
    fig, ax = plt.subplots(figsize=(8, 3.6))
    att[["Routine error share_%", "Large-payment error share_%"]].plot.barh(stacked=True, ax=ax,
                                                                           color=["#1b6ca8", "#f0b400"])
    ax.set_xlabel("% of total absolute hourly error (TEST)")
    ax.set_title("Where does the hourly error come from?", fontsize=10)
    ax.grid(alpha=.3, axis="x")
    return fig_to_b64(fig)


def html_table(df: pd.DataFrame, fmt="{:,.0f}", pct_cols=()) -> str:
    d = df.copy()
    for c in d.columns:
        if pd.api.types.is_numeric_dtype(d[c]):
            f = "{:.1f}" if c in pct_cols else fmt
            d[c] = d[c].map(lambda v: "" if pd.isna(v) else f.format(v))
    return d.to_html(classes="t", border=0, escape=False)




# ============================================================================ excel
def write_excel(path, res: pd.DataFrame, op, comp: pd.DataFrame, final: pd.DataFrame, outlook: pd.DataFrame,
                settings: pd.DataFrame, splits: dict):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter as L

    F, FB = Font(name="Arial", size=10), Font(name="Arial", size=10, bold=True, color="FFFFFF")
    FILL = PatternFill("solid", start_color="1F4E79")
    NUM, PCT = "#,##0;(#,##0);-", "0.0%"
    wb = Workbook()

    def header(ws, cols, row=1):
        for j, c in enumerate(cols, 1):
            cell = ws.cell(row=row, column=j, value=c)
            cell.font, cell.fill = FB, FILL
            cell.alignment = Alignment(horizontal="center", wrap_text=True)

    ws = wb.active
    ws.title = "Hourly_Test"
    cols = ["Hour_Start", "Date", "Hour", "Operating_Hour",
            "Credit_Actual", "Credit_Forecast", "Credit_P10", "Credit_P90", "Credit_Error", "Credit_AbsError", "Credit_InBand",
            "Debit_Actual", "Debit_Forecast", "Debit_P10", "Debit_P90", "Debit_Error", "Debit_AbsError", "Debit_InBand",
            "Net_Actual", "Net_Forecast", "Net_P10", "Net_P90", "Net_Error", "Net_AbsError",
            "Credit_Routine_Actual", "Credit_Routine_Forecast", "Credit_Large_Actual", "Credit_Large_Forecast",
            "Credit_Large_P90", "Credit_Scheduled_Actual", "Credit_Scheduled_Forecast",
            "Debit_Routine_Actual", "Debit_Routine_Forecast", "Debit_Large_Actual", "Debit_Large_Forecast",
            "Debit_Large_P90", "Debit_Scheduled_Actual", "Debit_Scheduled_Forecast"]
    header(ws, cols)
    for i, (ts, r) in enumerate(res.iterrows(), start=2):
        v = [ts.to_pydatetime(), ts.normalize().to_pydatetime(), ts.hour, int(ts.hour in op),
             r.credit_actual, f"=Z{i}+AB{i}+AE{i}", r.credit_p10, r.credit_p90, f"=F{i}-E{i}", f"=ABS(I{i})",
             f"=IF(AND(E{i}>=G{i},E{i}<=H{i}),1,0)",
             r.debit_actual, f"=AG{i}+AI{i}+AL{i}", r.debit_p10, r.debit_p90, f"=M{i}-L{i}", f"=ABS(P{i})",
             f"=IF(AND(L{i}>=N{i},L{i}<=O{i}),1,0)",
             f"=E{i}-L{i}", f"=F{i}-M{i}", r.net_p10, r.net_p90, f"=T{i}-S{i}", f"=ABS(W{i})",
             r.credit_routine_actual, r.credit_routine_forecast, r.credit_large_actual, r.credit_large_forecast,
             r.credit_large_p90, r.credit_sched_actual, r.credit_sched_forecast,
             r.debit_routine_actual, r.debit_routine_forecast, r.debit_large_actual, r.debit_large_forecast,
             r.debit_large_p90, r.debit_sched_actual, r.debit_sched_forecast]
        for j, x in enumerate(v, 1):
            c = ws.cell(row=i, column=j, value=float(x) if isinstance(x, (np.floating, np.integer)) else x)
            c.font = F
            c.number_format = ("yyyy-mm-dd hh:mm" if j == 1 else "yyyy-mm-dd" if j == 2 else
                               "0" if j in (3, 4, 11, 18) else NUM)
    last = len(res) + 1
    ws.freeze_panes = "E2"
    for j in range(1, len(cols) + 1):
        ws.column_dimensions[L(j)].width = 17 if j == 1 else 14

    sm = wb.create_sheet("Summary", 0)
    sm["A1"] = "60-min-ahead forecast - TEST period accuracy (operating hours only)"
    sm["A1"].font = Font(name="Arial", size=12, bold=True)
    sm["A2"] = (f"Test: {splits['test'][0]} to {splits['test'][1]}  |  Operating hours: "
                f"{', '.join(f'{h:02d}' for h in op)}")
    sm["A2"].font = F
    header(sm, ["Metric", "Credit", "Debit", "Net"], row=4)
    R = lambda c: f"Hourly_Test!${c}$2:${c}${last}"          # noqa: E731
    OP = R("D")
    rows = [
        ("TOTAL accuracy (1 - WAPE)", "=1-B6", "=1-C6", "=1-D6", PCT),
        ("TOTAL WAPE", f"=IFERROR(SUMPRODUCT({R('J')},{OP})/SUMPRODUCT({R('E')},{OP}),0)",
         f"=IFERROR(SUMPRODUCT({R('Q')},{OP})/SUMPRODUCT({R('L')},{OP}),0)",
         f"=IFERROR(SUMPRODUCT({R('X')},{OP})/SUMPRODUCT(ABS({R('S')}),{OP}),0)", PCT),
        ("TOTAL bias (forecast - actual) / actual", f"=IFERROR(SUMPRODUCT({R('I')},{OP})/SUMPRODUCT({R('E')},{OP}),0)",
         f"=IFERROR(SUMPRODUCT({R('P')},{OP})/SUMPRODUCT({R('L')},{OP}),0)",
         f"=IFERROR(SUMPRODUCT({R('W')},{OP})/SUMPRODUCT(ABS({R('S')}),{OP}),0)", PCT),
        ("P10-P90 coverage (target 80%)", f"=IFERROR(SUMPRODUCT({R('K')},{OP})/SUM({OP}),0)",
         f"=IFERROR(SUMPRODUCT({R('R')},{OP})/SUM({OP}),0)", "n/a", PCT),
        ("ROUTINE accuracy (1 - WAPE)",
         f"=IFERROR(1-SUMPRODUCT(ABS({R('Z')}-{R('Y')}),{OP})/SUMPRODUCT({R('Y')},{OP}),0)",
         f"=IFERROR(1-SUMPRODUCT(ABS({R('AG')}-{R('AF')}),{OP})/SUMPRODUCT({R('AF')},{OP}),0)", "n/a", PCT),
        ("Large payments - share of actual total", f"=IFERROR(SUM({R('AA')})/SUM({R('E')}),0)",
         f"=IFERROR(SUM({R('AH')})/SUM({R('L')}),0)", "n/a", PCT),
        ("Hours with a large payment", f"=COUNTIF({R('AA')},\">0\")", f"=COUNTIF({R('AH')},\">0\")", "n/a", NUM),
        ("Total actual (test)", f"=SUM({R('E')})", f"=SUM({R('L')})", f"=SUM({R('S')})", NUM),
        ("Total forecast (test)", f"=SUM({R('F')})", f"=SUM({R('M')})", f"=SUM({R('T')})", NUM),
    ]
    for i, (name, *fs, fmt) in enumerate(rows, start=5):
        sm.cell(row=i, column=1, value=name).font = F
        for j, f in enumerate(fs, start=2):
            c = sm.cell(row=i, column=j, value=f)
            c.font, c.number_format = F, fmt
    sm.cell(row=15, column=1, value="Rows 5-13 are live formulas on sheet Hourly_Test (Forecast = Routine + Large + "
            "Scheduled). Other sheets hold values computed by cashflow_forecast.py.").font = Font(name="Arial", italic=True, size=9)
    sm.column_dimensions["A"].width = 42
    for c in "BCD":
        sm.column_dimensions[c].width = 18

    wd = wb.create_sheet("Daily_Test")
    header(wd, ["Date", "Credit_Actual", "Credit_Forecast", "Credit_Accuracy", "Debit_Actual", "Debit_Forecast", "Debit_Accuracy"])
    for i, d in enumerate(sorted(res.index.normalize().unique()), start=2):
        wd.cell(row=i, column=1, value=d.to_pydatetime()).number_format = "yyyy-mm-dd"
        f = [f"=SUMIFS({R('E')},{R('B')},A{i})", f"=SUMIFS({R('F')},{R('B')},A{i})", f"=IFERROR(1-ABS(C{i}-B{i})/B{i},\"\")",
             f"=SUMIFS({R('L')},{R('B')},A{i})", f"=SUMIFS({R('M')},{R('B')},A{i})", f"=IFERROR(1-ABS(F{i}-E{i})/E{i},\"\")"]
        for j, x in enumerate(f, start=2):
            c = wd.cell(row=i, column=j, value=x)
            c.font, c.number_format = F, PCT if j in (4, 7) else NUM
    for j in range(1, 8):
        wd.column_dimensions[L(j)].width = 16

    def df_sheet(name, df, pct_cols=()):
        w = wb.create_sheet(name)
        d = df.reset_index()
        header(w, [str(c) for c in d.columns])
        for i, row in enumerate(d.itertuples(index=False), start=2):
            for j, x in enumerate(row, start=1):
                if isinstance(x, pd.Timestamp):
                    x = x.to_pydatetime()
                elif isinstance(x, (np.floating, np.integer)):
                    x = float(x)
                c = w.cell(row=i, column=j, value=x)
                c.font = F
                if isinstance(x, float):
                    c.number_format = "0.00" if str(d.columns[j - 1]) in pct_cols else NUM
                elif hasattr(x, "year"):
                    c.number_format = "yyyy-mm-dd hh:mm"
        for j in range(1, len(d.columns) + 1):
            w.column_dimensions[L(j)].width = 20

    df_sheet("Model_Comparison", comp, pct_cols=("Accuracy_%", "WAPE_%", "Bias_%", "sMAPE_%"))
    df_sheet("Final_Forecast", final)
    df_sheet("Outlook_24h", outlook)
    df_sheet("Settings", settings)
    wb.save(path)


# ============================================================================ html report
CSS = """
body{font-family:Segoe UI,Arial,sans-serif;max-width:1150px;margin:24px auto;padding:0 16px;color:#222;line-height:1.45}
h1{color:#1F4E79;margin-bottom:4px} h2{color:#1F4E79;border-bottom:2px solid #1F4E79;padding-bottom:4px;margin-top:34px}
.kpis{display:grid;grid-template-columns:repeat(auto-fit,minmax(200px,1fr));gap:12px;margin:16px 0}
.kpi{background:#f3f7fb;border-left:4px solid #1F4E79;padding:10px 14px;border-radius:4px}
.kpi b{font-size:1.5em;color:#1F4E79;display:block}
table.t{border-collapse:collapse;font-size:.85em;margin:8px 0;width:100%}
table.t th{background:#1F4E79;color:#fff;padding:5px 8px;text-align:right}
table.t td{padding:4px 8px;border-bottom:1px solid #e3e3e3;text-align:right}
table.t tr:nth-child(even){background:#f8f8f8}
.wrap{overflow-x:auto} img{max-width:100%} .ok{color:#2d8a4e;font-weight:600} .bad{color:#c0392b;font-weight:600}
.note{background:#fff8e1;border-left:4px solid #f0b400;padding:8px 12px;font-size:.92em}
code{background:#f1f1f1;padding:1px 4px;border-radius:3px}
"""


def write_report(path, c: dict):
    leak = "".join(
        f"<li>Cut {r['cut_hour']}: GBM features <span class='{'ok' if r['gbm_features_unchanged'] else 'bad'}'>"
        f"{'PASS' if r['gbm_features_unchanged'] else 'FAIL'}</span>, Chronos-2 context "
        f"<span class='{'ok' if r['fm_context_unchanged'] else 'bad'}'>"
        f"{'PASS' if r['fm_context_unchanged'] else 'FAIL'}</span></li>" for r in c["leak"])
    kpi = "".join(f"<div class='kpi'>{k}<b>{v}</b></div>" for k, v in c["kpis"].items())
    img = lambda k: f"<img src='data:image/png;base64,{c[k]}'>"   # noqa: E731
    html = f"""<!doctype html><html><head><meta charset='utf-8'>
<meta name="viewport" content="width=device-width, initial-scale=1">
<title>60-min Cash-Flow Forecast Report</title><style>{CSS}</style></head><body>
<h1>60-Minute-Ahead Cash-Flow Forecast</h1>
<div>Hybrid Chronos-2 (covariates) + Monte Carlo + Calendar features &nbsp;|&nbsp; Generated {c['generated']}
&nbsp;|&nbsp; Data {c['dq']['start']} &rarr; {c['dq']['end']}</div>
{c['demo_banner']}

<h2>1. Executive summary</h2>
<div class='kpis'>{kpi}</div>
<p>Accuracy = <b>1 &minus; WAPE</b> on the untouched TEST period, operating hours only.
WAPE = &Sigma;|forecast&minus;actual| / &Sigma;|actual| (MAPE is unusable on lumpy flows with near-zero hours).
<b>Routine accuracy</b> measures the part a model can realistically learn; <b>total accuracy</b> also carries the unpredictable large payments.</p>
<h3>Final forecast &mdash; next 60 minutes</h3><div class='wrap'>{c['final_html']}</div>

<h2>2. Data, quality and lumpy payments</h2>
<div class='wrap'>{c['dq_html']}</div>
<h3>Lumpy-payment split (thresholds learned on TRAIN only)</h3>
<div class='wrap'>{c['lumpy_html']}</div>
<p>Every minute amount is classified as <b>scheduled</b> (matches the scheduled-payments file), <b>large</b>
(&ge; the {c['lq']:.0%} quantile of TRAIN amounts, or a round multiple of {c['round_unit']:,.0f}) or <b>routine</b>.
Forecast = routine + large + scheduled.</p>

<h2>3. Method</h2>
<table class='t'><tr><th style='text-align:left'>Component</th><th style='text-align:left'>What it does</th><th style='text-align:left'>Status in this run</th></tr>
<tr><td style='text-align:left'>Chronos-2</td><td style='text-align:left'>Amazon zero-shot foundation model; credit &amp; debit routine flow forecast jointly (log scale); calendar features as future-known covariates, large-payment history as past-only covariates; context ensemble {c['contexts']} hours; output statistic chosen on validation ({c['fm_stat']})</td><td style='text-align:left'>{c['fm_status']}</td></tr>
<tr><td style='text-align:left'>Gradient boosting (optional)</td><td style='text-align:left'>{c['n_feat']} lag / rolling / intraday / large-payment / calendar features, walk-forward refit every {c['refit']} days</td><td style='text-align:left'>{c['gbm_engine']}</td></tr>
<tr><td style='text-align:left'>Monte Carlo &ndash; routine</td><td style='text-align:left'>{c['sims']} compound-Poisson simulations per hour from the same hour-of-week (last {c['K']} normal weeks) &times; calendar multipliers</td><td style='text-align:left'>active</td></tr>
<tr><td style='text-align:left'>Monte Carlo &ndash; large payments</td><td style='text-align:left'>probability &times; size: Poisson count by hour-of-day &amp; day type (last {c['KL']} weeks) &times; calendar factors, sizes bootstrapped</td><td style='text-align:left'>active</td></tr>
<tr><td style='text-align:left'>Monte Carlo &ndash; uncertainty</td><td style='text-align:left'>joint bootstrap of validation residuals (credit &amp; debit together) &rarr; P10/P50/P90</td><td style='text-align:left'>active</td></tr>
</table>
<h3>Hybrid weights and scales (learned on VALIDATION)</h3><div class='wrap'>{c['weights_html']}</div>
<h3>Calendar features</h3>
<p>hour, weekday, day-of-month, month, quarter, week-of-year (+ sin/cos), weekend, holiday, business-day-of-month,
business-days-to-month-end, first/last business day, quarter-end, year-end, mid-month, day before/after closed day,
days to/since holiday. Holidays supplied: <b>{c['n_hol']}</b>. Used by Chronos-2 (future-known covariates), gradient boosting if enabled (features)
and both Monte Carlo models (multipliers).</p>

<h3>Error-reduction steps (each fitted on VALIDATION)</h3>
<ol><li><b>Lumpy split</b> &mdash; large and scheduled payments removed before modelling, so Chronos-2 learns the smooth routine flow.</li>
<li><b>Log scale</b> &mdash; amounts from tens to hundreds of millions are modelled as log1p.</li>
<li><b>Context ensemble</b> &mdash; Chronos-2 run with several history lengths and averaged, reducing variance.</li>
<li><b>Output statistic selection</b> &mdash; Chronos-2's mean and 10th&ndash;90th quantiles are compared; the one with the lowest validation WAPE is used ({c['fm_stat']}).</li>
<li><b>Covariates</b> &mdash; calendar (holiday, month-end, weekday) known for the forecast hour; large-payment history and the other series (joint credit/debit forecast) as context.</li>
<li><b>Hybrid blend</b> &mdash; Chronos-2 and Monte Carlo weighted to minimise validation WAPE.</li>
<li><b>Hour-of-day calibration</b> &mdash; removes systematic over/under-forecasting per hour, shrunk towards 1 when data are few.</li>
<li><b>Large-payment scale</b> &mdash; how much of the Monte Carlo expected large payment to add, tuned on validation.</li></ol>
<h3>Ablation &mdash; routine flow, TEST (what each step contributes)</h3><div class='wrap'>{c['ablation']}</div>
<h3>Hour-of-day calibration factors</h3><div class='wrap'>{c['hscale_html']}</div>

<h2>4. Train / validation / test</h2><div class='wrap'>{c['split_html']}</div>

<h2>5. Data-leakage controls</h2>
<ul><li>Features and Chronos-2 contexts use data up to <i>h&minus;1&minus;latency</i> only (latency = {c['latency']} h). Calendar and scheduled payments are known in advance.</li>
<li>Large-payment thresholds, operating hours and calendar multipliers from TRAIN; hyper-parameters, blend weights, large-payment scale and residuals from VALIDATION; TEST scored once.</li>
<li>Walk-forward: each refit uses only hours already known at that moment. No shuffling, no future-fitted scalers.</li>
<li><b>Automated tamper test</b> (future randomised, past inputs must be identical):<ul>{leak}</ul></li></ul>

<h2>6. Accuracy &mdash; TEST (operating hours)</h2>
<h3>Total flow</h3><div class='wrap'>{c['comp_total']}</div>
<h3>Routine flow</h3><div class='wrap'>{c['comp_routine']}</div>
<h3>Large payments</h3><div class='wrap'>{c['large_html']}</div>
{img('img_attr')}
<p>The attribution shows how much of the remaining hourly error comes from large payments. That part can only be reduced
with information about payments before they happen (the <code>--schedule-file</code> input).</p>
{img('img_byhour')}
<h3>Validation results (used for model selection)</h3><div class='wrap'>{c['comp_val']}</div>

<h2>7. Hourly actual vs hourly forecast (TEST)</h2>
{img('img_hourly')}{img('img_routine')}{img('img_scatter')}
<h3>Last 48 operating hours</h3><div class='wrap'>{c['last48_html']}</div>
<p>Every test hour is in <code>forecast_results.xlsx</code> (sheet Hourly_Test, with live formulas) and <code>hourly_test_forecasts.csv</code>.</p>

<h2>8. Daily roll-up (TEST)</h2>{img('img_daily')}<div class='wrap'>{c['daily_html']}</div>

{"<h2>9. Gradient-boosting feature importance (credit, routine)</h2>" + img('img_imp') if c['img_imp'] else ""}

<h2>10. 24-hour outlook</h2>
<p class='note'>Only the first hour is the validated 60-minute-ahead forecast. Later hours have no fresh actuals and are indicative.</p>
<div class='wrap'>{c['outlook_html']}</div>

<h2>11. Limitations and next steps</h2><ul>
<li>Unscheduled large payments cannot be predicted from history; the Monte Carlo gives their probability and P90 instead.
Supplying known settlements, payroll, loan and tax payments via <code>--schedule-file</code> is the biggest single accuracy lever.</li>
<li>Retrain weekly; monitor rolling routine WAPE, total WAPE and P10&ndash;P90 coverage (target 80%).</li>
<li>Use your official holiday calendar (<code>--holiday-country</code> or <code>--holiday-file</code>).</li></ul>
</body></html>"""
    Path(path).write_text(html, encoding="utf-8")


# ============================================================================ main
def main():
    cfg = parse_args()
    out = Path(cfg.outdir)
    out.mkdir(parents=True, exist_ok=True)
    t0 = time.time()
    rng_note = []

    # ---- 1. data & split dates -------------------------------------------------------
    minute, dq = load_minute_data(cfg)
    sched = load_schedule(cfg)
    first, data_end = minute.ts.min().normalize(), minute.ts.max().floor("h")
    test_start = data_end.normalize() - pd.Timedelta(days=cfg.test_days - 1)
    val_start = test_start - pd.Timedelta(days=cfg.val_days)
    fit_start = first + cfg.warmup_hours * HOUR
    if val_start - fit_start < pd.Timedelta(days=60):
        raise ValueError("Need >= 60 days of TRAIN after the 4-week warm-up; reduce --val-days/--test-days.")

    # ---- 2. lumpy split (thresholds on TRAIN only) + hourly --------------------------
    m, thr = classify_minutes(minute, sched, val_start, cfg)
    H, _ = build_hourly(m, sched, cfg.outlook_hours)
    data_end = H.index.max()
    log(f"{dq['minute_rows']:,} minute rows -> {len(H):,} hours; large thresholds "
        + ", ".join(f"{s} {v:,.0f}" for s, v in thr.items()))
    ext = pd.date_range(H.index.min(), data_end + cfg.outlook_hours * HOUR, freq="h")
    Hx = H.reindex(ext)
    sched_h = pd.DataFrame(0.0, index=ext, columns=list(SERIES))
    if len(sched):
        sh = sched.assign(hour=sched.ts.dt.floor("h")).groupby("hour")[list(SERIES)].sum()
        sched_h.loc[sched_h.index.intersection(sh.index)] = sh.reindex(sched_h.index.intersection(sh.index)).to_numpy()
    hol = load_holidays(cfg, range(first.year, data_end.year + 2))
    cal = calendar_features(ext, hol)
    X = build_features(Hx, cal, sched_h, cfg)
    feats = list(X.columns)
    Yr = {s: np.log1p(Hx[f"{s}_r"]) for s in SERIES}

    leak = leakage_check(H, cal.loc[H.index], sched_h, cfg)
    passed = all(r["gbm_features_unchanged"] and r["fm_context_unchanged"] for r in leak)
    log(f"Leakage tamper test: {'PASS' if passed else 'FAIL'}")
    if not passed:
        raise RuntimeError("Leakage detected - aborting")

    idx = H.index
    train_mask = (idx >= fit_start) & (idx < val_start)
    val_idx, test_idx = idx[(idx >= val_start) & (idx < test_start)], idx[idx >= test_start]
    op = operating_hours(H, train_mask, cfg.op_hour_threshold)
    op_val, op_test = val_idx.hour.isin(op), test_idx.hour.isin(op)
    splits = {"train": (str(fit_start), str(val_start - HOUR)), "validation": (str(val_start), str(test_start - HOUR)),
              "test": (str(test_start), str(data_end))}
    log(f"Split train {splits['train']} | val {splits['validation']} | test {splits['test']}")

    fm, fm_status = load_fm(cfg)
    pools_r = build_pools(m, "_r")

    contexts = sorted(int(x) for x in cfg.fm_contexts.split(","))
    fm_stat = {}

    def fm_ensemble(positions, horizon=1):
        """Chronos-2 averaged (log space) over the context ensemble -> {series: (B, horizon, 10)}."""
        runs = []
        for C in contexts:
            log(f"  Chronos-2: {len(positions):,} origins x 2 series, context {C}h ...")
            runs.append(fm.predict(H, cal, positions, C, horizon))
        return {s: np.mean([r[s] for r in runs], axis=0) for s in SERIES}

    def routine_components(period_idx, pmask, seg_start, seg_end, fit_end, params=None, choose=False):
        """Routine-flow components for a period; every input strictly before its hour."""
        comps, best_p, cands = {s: {} for s in SERIES}, {}, {}
        mc = MonteCarloStructural(pools_r, H, cal, fit_end, cfg).forecast(period_idx)
        fmq = fm_ensemble(H.index.get_indexer(period_idx[pmask])) if fm is not None else None
        for s in SERIES:
            act = H[f"{s}_r"].loc[period_idx].to_numpy()
            comps[s]["MonteCarlo"] = mc[s]["p50"].to_numpy()
            if fm is not None:
                q = np.expm1(fmq[s][:, 0, :])                       # (B, 10)
                cands[s] = {}
                for j, name in enumerate(FM_STATS):
                    v = comps[s]["MonteCarlo"].copy()                 # non-operating hours -> MC (~0)
                    v[pmask] = q[:, j]
                    cands[s][name] = v
                if choose:                                            # validation: pick best statistic
                    errs = {k: wape(act[pmask], v[pmask]) for k, v in cands[s].items()}
                    fm_stat[s] = min(errs, key=errs.get)
                    log(f"  Chronos-2 {s}: best output statistic on validation = {fm_stat[s]} "
                        f"(WAPE {100 * errs[fm_stat[s]]:.1f}%, mean {100 * errs['mean']:.1f}%, q50 {100 * errs['q50']:.1f}%)")
                comps[s]["Chronos2"] = cands[s][fm_stat[s]]
            if cfg.use_gbm:
                if params is None:
                    best = np.inf
                    for p in param_grid(cfg.quick):
                        pr, _ = walk_forward(X, Yr[s], feats, fit_start, seg_start, seg_end, p, cfg)
                        e = wape(act[pmask], np.expm1(pr.to_numpy())[pmask])
                        log(f"  GBM {s} {p} -> routine val WAPE {100 * e:.1f}%")
                        if e < best:
                            best, best_p[s], g = e, p, pr
                else:
                    best_p[s] = params[s]
                    g, _ = walk_forward(X, Yr[s], feats, fit_start, seg_start, seg_end, params[s], cfg)
                comps[s]["GBM"] = np.expm1(g.to_numpy())
        return comps, best_p

    # ---- 3. VALIDATION: tune GBM, blend weights, large scale, residuals --------------
    log("VALIDATION ...")
    comps_val, best_p = routine_components(val_idx, op_val, val_start, test_start - HOUR, val_start, choose=True)
    weights, rscale, hscale, hyb_r_val, blend_nocal_val = {}, {}, {}, {}, {}
    for s in SERIES:
        act = H[f"{s}_r"].loc[val_idx].to_numpy()
        weights[s] = fit_blend(act, comps_val[s], op_val)
        raw = apply_blend(comps_val[s], weights[s])
        rscale[s] = fit_scale(act, raw, op_val, "wape")
        blend_nocal_val[s] = rscale[s] * raw
        hscale[s] = fit_hour_scale(act, blend_nocal_val[s], val_idx.hour, op_val, cfg.hour_calib_shrink)
        hyb_r_val[s] = apply_hour_scale(blend_nocal_val[s], val_idx.hour, hscale[s])
        log(f"  routine blend {s}: {weights[s]} x{rscale[s]:.2f}; hour calibration "
            f"{min(hscale[s].values(), default=1):.2f}-{max(hscale[s].values(), default=1):.2f}")
    lmc_val = LargePaymentMC(m, H, cal, val_start, cfg).forecast(val_idx)
    lscale, tot_val = {}, {}
    for s in SERIES:
        lscale[s] = fit_large_scale(H[s].loc[val_idx], hyb_r_val[s], lmc_val[s]["mean"], sched_h[s].loc[val_idx],
                                    op_val, cfg.objective)
        tot_val[s] = hyb_r_val[s] + lscale[s] * lmc_val[s]["mean"].to_numpy() + sched_h[s].loc[val_idx].to_numpy()
        log(f"  large-payment scale {s}: x{lscale[s]:.2f} ({cfg.objective})")
    rmc = ResidualMonteCarlo(val_idx, {s: H[s].loc[val_idx] for s in SERIES}, tot_val, cfg)

    # ---- 4. TEST (scored once) ---------------------------------------------------------
    log("TEST ...")
    comps_test, _ = routine_components(test_idx, op_test, test_start, data_end, test_start, params=best_p)
    lmc_test = LargePaymentMC(m, H, cal, test_start, cfg).forecast(test_idx)
    res = pd.DataFrame(index=test_idx)
    tot_test, hyb_r_test = {}, {}
    for s in SERIES:
        blend_nocal = rscale[s] * apply_blend(comps_test[s], weights[s])
        comps_test[s]["Blend (no hour calibration)"] = blend_nocal
        hyb_r_test[s] = apply_hour_scale(blend_nocal, test_idx.hour, hscale[s])
        L = lscale[s] * lmc_test[s]["mean"].to_numpy()
        sc = sched_h[s].loc[test_idx].to_numpy()
        tot_test[s] = hyb_r_test[s] + L + sc
        res[f"{s}_actual"] = H[s].loc[test_idx]
        res[f"{s}_forecast"] = tot_test[s]
        res[f"{s}_routine_actual"] = H[f"{s}_r"].loc[test_idx]
        res[f"{s}_routine_forecast"] = hyb_r_test[s]
        res[f"{s}_large_actual"] = H[f"{s}_L"].loc[test_idx]
        res[f"{s}_large_forecast"] = L
        res[f"{s}_large_p90"] = lmc_test[s]["p90"].to_numpy()
        res[f"{s}_large_prob"] = lmc_test[s]["prob"].to_numpy()
        res[f"{s}_sched_actual"] = H[f"{s}_sched"].loc[test_idx]
        res[f"{s}_sched_forecast"] = sc
    res["net_actual"] = res.credit_actual - res.debit_actual
    res["net_forecast"] = res.credit_forecast - res.debit_forecast
    unc = rmc.simulate(test_idx, tot_test)
    for k in (*SERIES, "net"):
        for q in ("p10", "p50", "p90"):
            res[f"{k}_{q}"] = unc[k][q]
    res.index.name = "hour_start"

    # ---- 5. comparison tables ----------------------------------------------------------
    kk = 1 + cfg.latency_hours

    def table(period, pidx, pmask, comps, hyb_r, lmc):
        rows_t, rows_r = [], []
        models_r = {"Naive (t-1h)": {s: H[f"{s}_r"].shift(kk).loc[pidx].fillna(0).to_numpy() for s in SERIES},
                    "Seasonal naive (t-1w)": {s: H[f"{s}_r"].shift(168).loc[pidx].fillna(0).to_numpy() for s in SERIES}}
        for name in comps["credit"]:
            models_r[name] = {s: comps[s][name] for s in SERIES}
        models_r["HYBRID"] = hyb_r
        for name, fc in models_r.items():
            tot = {s: fc[s] + lscale[s] * lmc[s]["mean"].to_numpy() + sched_h[s].loc[pidx].to_numpy() for s in SERIES}
            if name.startswith(("Naive", "Seasonal")):   # classic baselines on the raw total
                tot = {s: H[s].shift(kk if name.startswith("Naive") else 168).loc[pidx].fillna(0).to_numpy() for s in SERIES}
            for tgt in ("credit", "debit", "net"):
                if tgt == "net":
                    a = (H.credit - H.debit).loc[pidx].to_numpy()
                    f = tot["credit"] - tot["debit"]
                    ar = (H.credit_r - H.debit_r).loc[pidx].to_numpy()
                    fr = fc["credit"] - fc["debit"]
                else:
                    a, f = H[tgt].loc[pidx].to_numpy(), tot[tgt]
                    ar, fr = H[f"{tgt}_r"].loc[pidx].to_numpy(), fc[tgt]
                label = name if name.startswith(("Naive", "Seasonal", "HYBRID")) else f"{name} (+large MC)"
                rows_t.append({"Target": tgt, "Model": label, **metrics(a[pmask], np.asarray(f)[pmask])})
                rows_r.append({"Target": tgt, "Model": name, **metrics(ar[pmask], np.asarray(fr)[pmask])})
        return (pd.DataFrame(rows_t).set_index(["Target", "Model"]),
                pd.DataFrame(rows_r).set_index(["Target", "Model"]))

    comp_total, comp_routine = table("Test", test_idx, op_test, comps_test, hyb_r_test, lmc_test)
    for s in SERIES:
        comps_val[s]["Blend (no hour calibration)"] = blend_nocal_val[s]
    comp_total_v, _ = table("Validation", val_idx, op_val, comps_val, hyb_r_val, lmc_val)
    opres = res[op_test]
    large_rows, att_rows = [], []
    for s in SERIES:
        la, lf = opres[f"{s}_large_actual"], opres[f"{s}_large_forecast"]
        hit = la > 0
        large_rows.append({"Series": s, "Hours with large payment": int(hit.sum()),
                           "Large share of total value_%": 100 * la.sum() / max(opres[f"{s}_actual"].sum(), 1),
                           "Mean predicted probability_%": 100 * opres[f"{s}_large_prob"].mean(),
                           "Actual frequency_%": 100 * hit.mean(),
                           "Actual <= MC P90_%": 100 * (la <= opres[f"{s}_large_p90"]).mean(),
                           "WAPE forecast=0_%": 100 * wape(la, np.zeros(len(la))),
                           "WAPE large MC_%": 100 * wape(la, lf)})
        er = (opres[f"{s}_routine_forecast"] - opres[f"{s}_routine_actual"]).abs().sum()
        el = (lf + opres[f"{s}_sched_forecast"] - la - opres[f"{s}_sched_actual"]).abs().sum()
        att_rows.append({"Series": s, "Routine error share_%": 100 * er / max(er + el, 1),
                         "Large-payment error share_%": 100 * el / max(er + el, 1)})
    large_df = pd.DataFrame(large_rows).set_index("Series")
    att = pd.DataFrame(att_rows).set_index("Series")
    cov = {s: float(((opres[f"{s}_actual"] >= opres[f"{s}_p10"]) & (opres[f"{s}_actual"] <= opres[f"{s}_p90"])).mean())
           for s in (*SERIES, "net")}
    daily = res.groupby(res.index.normalize()).sum()
    daily_m = pd.DataFrame({s: metrics(daily[f"{s}_actual"], daily[f"{s}_forecast"]) for s in (*SERIES, "net")}).T

    # ---- 6. FINAL: next 60 minutes + 24h outlook -------------------------------------
    log("Final refit on all data ...")
    nxt = data_end + HOUR
    out_idx = pd.date_range(nxt, periods=cfg.outlook_hours, freq="h")
    all_mask = (X.index >= fit_start) & (X.index <= data_end)
    mc_f = MonteCarloStructural(pools_r, H, cal, nxt, cfg).forecast(out_idx, cutoff=nxt)
    lmc_f = LargePaymentMC(m, H, cal, nxt, cfg).forecast(out_idx, cutoff=nxt)
    final_models, fin, outlook = {}, {}, pd.DataFrame(index=out_idx)
    fm_final = fm_ensemble([len(H)], horizon=cfg.outlook_hours) if fm is not None else None
    for s in SERIES:
        c1 = {"MonteCarlo": np.array([mc_f[s]["p50"].iloc[0]])}
        if cfg.use_gbm:
            final_models[s] = make_model(best_p[s], cfg).fit(X.loc[all_mask, feats], Yr[s][all_mask])
            c1["GBM"] = np.array([np.expm1(max(final_models[s].predict(X.loc[[nxt], feats])[0], 0))])
        fm24 = None
        if fm is not None:
            j = FM_STATS.index(fm_stat[s])
            fm24 = np.expm1(fm_final[s][0, :, j])
            fm24 = apply_hour_scale(fm24, out_idx.hour, hscale[s])
            c1["Chronos2"] = fm24[:1] / hscale[s].get(nxt.hour, 1.0)
        routine1 = apply_hour_scale(rscale[s] * apply_blend(c1, weights[s]), [nxt.hour], hscale[s])
        L = lscale[s] * lmc_f[s]["mean"].to_numpy()
        sc = sched_h[s].loc[out_idx].to_numpy()
        fin[s] = routine1 + L[:1] + sc[:1]
        fin[f"{s}_parts"] = (routine1[0], L[0], sc[0])
        r24 = fm24 if fm24 is not None else mc_f[s]["p50"].to_numpy()
        outlook[f"{s}_routine"] = r24
        outlook[f"{s}_large_expected"] = L
        outlook[f"{s}_large_p90"] = lmc_f[s]["p90"].to_numpy()
        outlook[f"{s}_scheduled"] = sc
        outlook[f"{s}_total"] = r24 + L + sc
    outlook.iloc[0, outlook.columns.get_indexer([f"{s}_routine" for s in SERIES])] = [fin[f"{s}_parts"][0] for s in SERIES]
    for s in SERIES:
        outlook[f"{s}_total"].iloc[0] = fin[s][0]
    outlook["net_total"] = outlook.credit_total - outlook.debit_total
    outlook.index.name = "hour_start"
    unc_f = rmc.simulate(pd.DatetimeIndex([nxt]), {s: fin[s] for s in SERIES})
    final = pd.DataFrame({
        "forecast_window": [f"{nxt:%Y-%m-%d %H:%M} - {nxt + pd.Timedelta(minutes=59):%H:%M}"],
        **{f"{s}_forecast": fin[s] for s in SERIES}, "net_forecast": fin["credit"] - fin["debit"],
        **{f"{s}_{part}": [fin[f"{s}_parts"][j]] for s in SERIES for j, part in enumerate(("routine", "large_expected", "scheduled"))},
        **{f"{s}_large_p90": [lmc_f[s]["p90"].iloc[0]] for s in SERIES},
        **{f"{k}_{q}": unc_f[k][q].to_numpy() for k in (*SERIES, "net") for q in ("p10", "p50", "p90")}},
        index=pd.Index([nxt], name="hour_start"))

    # ---- 7. importance ------------------------------------------------------------------
    imp = None
    if not cfg.use_gbm:
        pass
    elif gbm_engine(cfg) in ("lightgbm", "xgboost"):
        imp = pd.Series(final_models["credit"].feature_importances_, index=feats)
    else:
        samp = X.loc[test_idx[op_test], feats].tail(1500)
        pi = permutation_importance(final_models["credit"], samp, Yr["credit"].loc[samp.index], n_repeats=3,
                                    random_state=cfg.seed, scoring="neg_mean_absolute_error")
        imp = pd.Series(pi.importances_mean, index=feats)

    # ---- 8. outputs ---------------------------------------------------------------------
    res.to_csv(out / "hourly_test_forecasts.csv")
    final.to_csv(out / "final_forecast.csv")
    outlook.to_csv(out / "outlook_24h.csv")
    comp_total.to_csv(out / "model_comparison_total.csv")
    comp_routine.to_csv(out / "model_comparison_routine.csv")
    wdf = pd.DataFrame(weights).fillna(0)
    wdf.loc["Routine calibration x"] = pd.Series(rscale)
    wdf.loc["Large-payment scale x"] = pd.Series(lscale)
    wdf.loc["Large threshold (TRAIN)"] = pd.Series(thr)
    if fm_stat:
        wdf.loc["Chronos-2 output statistic"] = pd.Series(fm_stat)
    settings = wdf.copy()
    settings.index.name = "setting"
    write_excel(out / "forecast_results.xlsx", res, op, comp_total.round(4), final, outlook,
                settings, splits)

    ct = lambda t, mdl: comp_total.loc[(t, mdl), "Accuracy_%"]       # noqa: E731
    cr = lambda t: comp_routine.loc[(t, "HYBRID"), "Accuracy_%"]     # noqa: E731
    kpis = {f"Routine hourly accuracy &ndash; {s}": f"{cr(s):.1f}%" for s in SERIES}
    kpis.update({f"Total hourly accuracy &ndash; {s}": f"{ct(s, 'HYBRID'):.1f}%" for s in SERIES})
    kpis["Total daily accuracy &ndash; credit / debit"] = f"{daily_m.loc['credit', 'Accuracy_%']:.1f}% / {daily_m.loc['debit', 'Accuracy_%']:.1f}%"
    kpis["P10&ndash;P90 coverage credit / debit"] = f"{cov['credit']:.0%} / {cov['debit']:.0%}"
    last48 = opres.tail(48)[["credit_actual", "credit_forecast", "credit_p10", "credit_p90",
                             "debit_actual", "debit_forecast", "net_actual", "net_forecast"]]
    last48.index = last48.index.strftime("%Y-%m-%d %H:%M")
    split_df = pd.DataFrame({"From": {k: v[0] for k, v in splits.items()}, "To": {k: v[1] for k, v in splits.items()},
                             "Hours": {"train": int(train_mask.sum()), "validation": len(val_idx), "test": len(test_idx)},
                             "Used for": {"train": "GBM fit, large thresholds, calendar multipliers, operating hours",
                                          "validation": "GBM hyper-parameters, blend weights, scales, residual distribution",
                                          "test": "scored once (walk-forward)"}})
    lumpy = pd.DataFrame({s: {"Large threshold (TRAIN)": f"{thr[s]:,.0f}",
                              "Large payments (all data)": f"{int(H[f'{s}_Ln'].sum()):,}",
                              "Large share of value": f"{100 * H[f'{s}_L'].sum() / max(H[s].sum(), 1):.1f}%",
                              "Round-amount share of value": f"{100 * H[f'{s}_round'].sum() / max(H[s].sum(), 1):.1f}%",
                              "Scheduled share of value": f"{100 * H[f'{s}_sched'].sum() / max(H[s].sum(), 1):.1f}%"}
                          for s in SERIES})
    dq_df = pd.DataFrame.from_dict({k: (f"{v:,.0f}" if isinstance(v, float) else v) for k, v in dq.items()},
                                   orient="index", columns=["value"])
    fin_show = final.T.rename(columns=lambda c: "value")
    num = fin_show.index != "forecast_window"
    fin_show.loc[num, "value"] = fin_show.loc[num, "value"].map(lambda v: f"{float(v):,.0f}")
    pct = ("Accuracy_%", "WAPE_%", "Bias_%", "sMAPE_%")
    demo = ""
    ctx = dict(
        generated=time.strftime("%Y-%m-%d %H:%M"), dq=dq, demo_banner=demo, kpis=kpis,
        final_html=fin_show.to_html(classes="t", border=0), dq_html=dq_df.to_html(classes="t", border=0),
        lumpy_html=lumpy.to_html(classes="t", border=0), lq=cfg.large_quantile, round_unit=cfg.round_unit,
        fm_status=(fm_status + (f" &mdash; {fm.cov_status}" if fm is not None else "")), fm_ctx=cfg.fm_context_hours,
        gbm_engine=(gbm_engine(cfg) if cfg.use_gbm else "not used (add --use-gbm)"), n_feat=len(feats), refit=cfg.refit_every_days, sims=cfg.mc_sims,
        K=cfg.mc_lookback_weeks, KL=cfg.large_lookback_weeks, n_hol=len(hol), latency=cfg.latency_hours,
        weights_html=wdf.to_html(classes="t", border=0, float_format=lambda v: f"{v:,.2f}"),
        split_html=split_df.to_html(classes="t", border=0), leak=leak,
        comp_total=html_table(comp_total, pct_cols=pct), comp_routine=html_table(comp_routine, pct_cols=pct),
        comp_val=html_table(comp_total_v, pct_cols=pct),
        large_html=html_table(large_df, fmt="{:,.1f}", pct_cols=()), img_attr=chart_attribution(att),
        img_byhour=chart_by_hour(res, op, {s: H[s].shift(168) for s in SERIES}), img_hourly=chart_hourly(res),
        img_routine=chart_routine(res), img_scatter=chart_scatter(res, op),
        last48_html=html_table(last48), img_daily=chart_daily(daily), daily_html=html_table(daily_m, pct_cols=pct),
        img_imp=chart_importance(imp) if imp is not None else None,
        fm_stat=", ".join(f"{s}: {v}" for s, v in fm_stat.items()) or "n/a", contexts=cfg.fm_contexts,
        hscale_html=pd.DataFrame(hscale).sort_index().to_html(classes="t", border=0, float_format=lambda v: f"{v:.2f}"),
        ablation=html_table(comp_routine.loc[(slice(None), [m_ for m_ in comp_routine.index.get_level_values(1).unique()
                                                           if not m_.startswith(("Naive", "Seasonal"))]), ["Accuracy_%", "WAPE_%", "Bias_%"]],
                            pct_cols=pct),
        outlook_html=html_table(outlook.set_index(outlook.index.strftime("%Y-%m-%d %H:%M"))))
    write_report(out / "forecast_report.html", ctx)
    (out / "run_config.json").write_text(json.dumps({**asdict(cfg), "chronos2": fm_status, "gbm_engine": gbm_engine(cfg),
                                                     "best_params": best_p, "blend_weights": weights, "chronos2_statistic": fm_stat,
                                                     "hour_calibration": hscale,
                                                     "routine_scale": rscale, "large_scale": lscale,
                                                     "large_thresholds": thr, "operating_hours": op,
                                                     "splits": splits, "coverage": cov}, indent=2, default=str))
    log(f"Done in {time.time() - t0:.0f}s -> {out.resolve()}")
    print(comp_total.round(1).to_string())
    print(comp_routine.loc[(slice(None), "HYBRID"), :].round(1).to_string())
    print(large_df.round(1).to_string())
    print(att.round(1).to_string())
    print("\nFINAL next 60 min:\n", final.T.to_string())


if __name__ == "__main__":
    main()
